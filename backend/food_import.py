"""Validate MEXT 2020 corrected workbook before atomic structured import.
Run without --import-db first; all validation also runs before any database writes.
"""
import argparse
import hashlib
import json
import re
from collections import Counter
from pathlib import Path
from decimal import Decimal
import openpyxl
import pymupdf
from psycopg.types.json import Jsonb
from backend.knowledge import connect

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / 'data/food2020'
PDF = Path('/home/aitd/下载/日本食品標準成分表.pdf')
URL = 'https://www.mext.go.jp/content/20201225-mxt_kagsei-mext_01110_012.xlsx'

def parse_value(value):
    raw = '' if value is None else str(value)
    text = raw.strip()
    estimated = text.startswith('(') and text.endswith(')')
    inner = text[1:-1] if estimated else text
    footnote = '†' if inner.endswith('†') else None
    if footnote:
        inner = inner[:-1]
    flags = {'estimated': estimated, 'footnote': footnote}
    specials = {'': 'blank', '-': 'not_reported', 'Tr': 'trace', '*': 'asterisk_unresolved'}
    if inner in specials:
        return {'raw': raw, 'value': None, 'status': specials[inner], **flags}
    if not re.fullmatch(r'\d+(?:\.\d+)?', inner):
        raise ValueError(f'Unrecognized nutrient value: {raw!r}')
    return {'raw': raw, 'value': str(Decimal(inner)), 'status': 'numeric', **flags}

def merged_value(sheet, row, col):
    value = sheet.cell(row, col).value
    if value is not None:
        return str(value)
    for area in sheet.merged_cells.ranges:
        if area.min_row <= row <= area.max_row and area.min_col <= col <= area.max_col:
            value = sheet.cell(area.min_row, area.min_col).value
            return '' if value is None else str(value)
    return ''

def validate():
    # Remove stale successful staging before validating a new run.
    for path in [DATA/'validated.json', ROOT/'data/reports/food2020_validation.json']:
        path.unlink(missing_ok=True)
    workbook = openpyxl.load_workbook(DATA/'012.xlsx', data_only=True)
    sheet = workbook['表全体']
    assert sheet['BJ1'].value == '更新日：2021年12月28日'
    nutrients = []
    for cell in sheet[12]:
        if cell.column < 5 or not cell.value:
            continue
        labels = []
        for r in range(3, 11):
            label = re.sub(r'\s+', '', merged_value(sheet, r, cell.column))
            if label and label not in labels:
                labels.append(label)
        unit_raw = merged_value(sheet, 11, cell.column)
        units = re.findall(r'kcal|kJ|mg|μg|g|%', unit_raw)
        assert len(units) == 1, (cell.coordinate, unit_raw)
        nutrients.append(dict(code=str(cell.value).strip(), column=cell.column_letter,
            column_number=cell.column, name_ja='/'.join(labels), unit=units[0],
            basis='whole_food_discard_percent' if cell.value == 'REFUSE' else 'per_100g_edible'))
    assert len({x['code'] for x in nutrients}) == len(nutrients) == 54
    foods = []
    for row in sheet.iter_rows(min_row=13):
        code = str(row[1].value)
        assert re.fullmatch(r'\d{5}', code)
        foods.append(dict(code=code, group=str(row[0].value), index=row[2].value,
            name=row[3].value, excel_sheet=sheet.title, excel_row=row[0].row,
            notes=row[61].value or '', energy_flags={'O': row[14].value or '', 'R': row[17].value or ''},
            nutrients={n['code']:parse_value(row[n['column_number']-1].value) for n in nutrients}, pdf_references=[]))
    assert len(foods) == len({f['code'] for f in foods}) == 2478
    by_code = {f['code']:f for f in foods}
    # Independently compare every data cell against the 18 food-group sheets.
    compared = 0
    for group in workbook.worksheets[1:]:
        for row in group.iter_rows(min_row=13):
            food = by_code[str(row[1].value)]
            source_row = sheet[food['excel_row']]
            assert [c.value for c in row[:62]] == [c.value for c in source_row[:62]], (group.title, row[0].row)
            compared += 1
    assert compared == 2478
    # PDF references are related context, never a fabricated nutrient-table page.
    pdf = pymupdf.open(PDF)
    energy_checked = 0
    for page_num, page in enumerate(pdf, 1):
        text = page.get_text()
        codes = set(re.findall(r'(?<!\d)\d{5}(?!\d)', text)) & by_code.keys()
        kind = ('energy_comparison' if 351 <= page_num <= 425 else
                'name_index' if 442 <= page_num <= 471 else
                'food_description' if 73 <= page_num <= 350 else 'other_context')
        printed = page_num + 176 if 73 <= page_num <= 427 else page_num + 181 if page_num >= 442 else None
        for code in codes:
            by_code[code]['pdf_references'].append(dict(pdf_page=page_num, printed_page=printed, kind=kind,
                verifies_nutrient_values=False))
        if kind == 'energy_comparison':
            words = page.get_text('words')
            for word in words:
                code = word[4]
                if not (60 < word[0] < 86 and re.fullmatch(r'\d{5}', code) and code in by_code):
                    continue
                matches = [w[4] for w in words if 290 < w[0] < 322 and abs(w[1]-word[1]) < 1
                           and re.fullmatch(r'\d+(?:\.\d+)?', w[4])]
                assert len(matches)==1, (page_num,code,matches)
                food=by_code[code]
                if 'pdf_energy_comparison' in food:
                    assert food['pdf_energy_comparison']['raw_kcal']==matches[0], code
                    continue
                food['pdf_energy_comparison']={'pdf_page':page_num,'raw_kcal':matches[0],
                    'matches_excel':Decimal(matches[0])==Decimal(food['nutrients']['ENERC_KCAL']['value'])}
                energy_checked+=1
    assert energy_checked == 2478
    errata_book=openpyxl.load_workbook(DATA/'101.xlsx',data_only=True)
    errata=errata_book.worksheets[0]
    errata_rows=list(errata.values)
    for f in foods:
        f['errata_entries']=[{'sheet':errata.title,'row':i,'item':str(row[5]),'before':str(row[6]),
            'after':str(row[7]),'revision':str(row[9])} for i,row in enumerate(errata_rows,1)
            if str(row[2])==f['code']]
    for sheet_name, energy_col in [('別添1',5),('本表',6)]:
        for i,row in enumerate(errata_book[sheet_name].values,1):
            code=str(row[1])
            if code in by_code and row[0]=='正':
                by_code[code]['errata_entries'].append({'sheet':sheet_name,'row':i,
                    'item':'energy_kcal','after':str(row[energy_col]),'revision':'official_corrected_row'})
    for f in foods:
        f['validation_status']='ready'
        if not f['pdf_energy_comparison']['matches_excel']:
            candidates=[e for e in f['errata_entries'] if 'kcal' in e['item']]
            expected=Decimal(f['nutrients']['ENERC_KCAL']['value'])
            matching=[e for e in candidates if re.fullmatch(r'\d+(?:\.\d+)?',e['after']) and Decimal(e['after'])==expected]
            if not matching:
                f['validation_status']='quarantined_pdf_energy_difference'
            f['pdf_energy_comparison']['correction_evidence']=matching
    assert all(any(r['kind']=='energy_comparison' for r in f['pdf_references']) for f in foods)
    # Fixed pilot values and corrected erratum, catches column shifts/unit mistakes.
    pilots = {'01088': {'ENERC_KCAL':'156','NA':'1','K':'29','PROT-':'2.5'},
              '06267': {'NA':'16','K':'690'}, '17007': {'NA':'5700','K':'390'}, '01026':{'ENERC':'1050'}}
    for code, expected in pilots.items():
        for nutrient,value in expected.items():
            assert Decimal(by_code[code]['nutrients'][nutrient]['value']) == Decimal(value)
    assert {n['code']:n['unit'] for n in nutrients}['NA']=='mg'
    assert {n['code']:n['unit'] for n in nutrients}['NACL_EQ']=='g'
    source = dict(edition='2020（八訂）, workbook updated 2021-12-28', url=URL,
        sha256=hashlib.sha256((DATA/'012.xlsx').read_bytes()).hexdigest(),
        pdf_path=str(PDF),pdf_sha256=hashlib.sha256(PDF.read_bytes()).hexdigest(),
        errata_sha256=hashlib.sha256((DATA/'101.xlsx').read_bytes()).hexdigest(),
        pdf_scope='Official PDF omits chapter 2; references support identity/context, not all nutrient values.')
    eligible=[f for f in foods if f['validation_status']=='ready']
    report=dict(eligible_foods=len(eligible), quarantined_foods=[f['code'] for f in foods if f['validation_status']!='ready'], source=source, foods=len(foods),fields=len(nutrients),values=len(foods)*len(nutrients),
        group_sheet_rows_verified=compared, pdf_energy_rows_parsed=energy_checked,
        pdf_energy_matches=sum(f.get('pdf_energy_comparison',{}).get('matches_excel',False) for f in foods),
        pdf_energy_differences=[{'code':f['code'],**f['pdf_energy_comparison']} for f in foods
            if 'pdf_energy_comparison' in f and not f['pdf_energy_comparison']['matches_excel']],
        statuses=dict(Counter(v['status'] for f in foods for v in f['nutrients'].values())),
        pilots=[by_code[c] for c in pilots], validation='passed_for_eligible_foods_with_quarantine')
    (DATA/'validated.json').write_text(json.dumps(dict(source=source,nutrients=nutrients,foods=foods),ensure_ascii=False))
    (ROOT/'data/reports/food2020_validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
    return source,nutrients,eligible,report

def import_database(source,nutrients,foods):
    with connect() as conn:
        conn.execute("SELECT pg_advisory_xact_lock(20202020)")
        conn.execute((ROOT/'backend/sql/002_foods.sql').read_text())
        existing=conn.execute('SELECT id FROM food_datasets WHERE sha256=%s',(source['sha256'],)).fetchone()
        if existing:
            raise ValueError('This source is already imported; unchanged dataset left intact')
        dataset=conn.execute('INSERT INTO food_datasets(sha256,metadata) VALUES (%s,%s) RETURNING id',
            (source['sha256'],Jsonb(source))).fetchone()['id']
        with conn.cursor() as cur:
            cur.executemany('INSERT INTO food_nutrients(dataset_id,code,metadata) VALUES (%s,%s,%s)',
                [(dataset,n['code'],Jsonb(n)) for n in nutrients])
            cur.executemany('INSERT INTO foods(dataset_id,code,name_ja,food_group,source) VALUES (%s,%s,%s,%s,%s)',
                [(dataset,f['code'],f['name'],f['group'],Jsonb({k:v for k,v in f.items() if k!='nutrients'})) for f in foods])
            with cur.copy('COPY food_values(dataset_id,food_code,nutrient_code,value,raw,status,estimated,footnote) FROM STDIN') as copy:
                for f in foods:
                    for code,v in f['nutrients'].items():
                        copy.write_row((dataset,f['code'],code,v['value'],v['raw'],v['status'],v['estimated'],v['footnote']))
        assert conn.execute('SELECT count(*) n FROM food_values WHERE dataset_id=%s',(dataset,)).fetchone()['n']==len(foods)*len(nutrients)
    return dataset

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--import-db',action='store_true');args=parser.parse_args()
    source,nutrients,foods,report=validate()
    print(json.dumps({k:v for k,v in report.items() if k not in ['pilots','pdf_energy_differences','source']},ensure_ascii=False))
    if args.import_db:print('Imported dataset',import_database(source,nutrients,foods))
